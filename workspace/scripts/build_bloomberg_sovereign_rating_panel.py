#!/usr/bin/env python3
"""Build beginning-of-year sovereign rating panels from Bloomberg rating-change exports.

The input archive is a static-value Bloomberg Excel export, not a direct country-year
panel. This script deduplicates overlapping rating-change rows, groups same-day events,
and infers beginning-of-year ratings from event histories.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import statistics
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_INPUT_ZIP = Path(
    "sources/ratings/bloomberg_static_rating_changes_2026-05-28/raw/"
    "bloomberg_rating_changes_static_values_2026-05-28.zip"
)
DEFAULT_OUTPUT_DIR = Path("data-raw/ratings/bloomberg_sovereign_rating_panel_2026-05-28")
DEFAULT_COUNTRY_METADATA = Path("data-raw/world_bank_countries.json")

YEARS = list(range(2000, 2026))

AGENCY_MAP = {
    "穆迪": "Moodys",
    "惠誉": "Fitch",
    "标普": "SP",
}

RATING_TYPE_MAP = {
    "发行人地方货币评级": "issuer_local_currency",
    "发行人外币评级": "issuer_foreign_currency",
    "发行人评级": "issuer_rating",
    "长期发行人违约评级": "long_term_issuer_default",
    "长期外币发行人信用": "long_term_foreign_currency_issuer_credit",
    "长期本币发行人信用": "long_term_local_currency_issuer_credit",
    "长期本币发行人违约": "long_term_local_currency_issuer_default",
}

PREFERRED_PAIRS = {
    ("Moodys", "issuer_foreign_currency"),
    ("Fitch", "long_term_issuer_default"),
    ("SP", "long_term_foreign_currency_issuer_credit"),
}

WITHDRAWN_OR_NOT_RATED = {"WD", "WR", "NR", "N.R.", "N/A", "NA", "--", ""}

EXPECTED_HEADERS = [
    "公司名称",
    "日期",
    "评级类型",
    "机构",
    "目前评级",
    "上次评级",
    "国家/地区",
    "行业类型",
    "证券名称",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-zip", type=Path, default=DEFAULT_INPUT_ZIP)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--country-metadata", type=Path, default=DEFAULT_COUNTRY_METADATA)
    return parser.parse_args()


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        value = value.strip()
        for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                continue
    return None


def clean_rating(value: Any) -> str | None:
    if value is None:
        return None
    rating = str(value).strip()
    if not rating:
        return None
    rating = re.sub(r"\s*\*[+-]?\s*$", "", rating).strip()
    if re.match(r"^(Aaa|Aa[123]|A[123]|Baa[123]|Ba[123]|B[123]|Caa[123]|Ca|C)u$", rating):
        rating = rating[:-1]
    if re.match(r"^(AAA|AA[+-]?|A[+-]?|BBB[+-]?|BB[+-]?|B[+-]?|CCC[+-]?|CC|C|RD|SD|D)u$", rating):
        rating = rating[:-1]
    return rating


def is_active_rating(value: str | None) -> bool:
    return value is not None and value not in WITHDRAWN_OR_NOT_RATED


def mode_or_none(values: list[Any]) -> Any:
    values = [v for v in values if v is not None and v != ""]
    if not values:
        return None
    counts = Counter(values)
    return sorted(counts.items(), key=lambda item: (-item[1], str(item[0])))[0][0]


def collapse_values(values: list[Any]) -> str:
    cleaned = sorted({str(v) for v in values if v is not None and str(v) != ""})
    return "; ".join(cleaned)


def read_country_metadata(path: Path) -> dict[str, dict[str, str]]:
    raw = json.loads(path.read_text())
    countries = raw[1]
    mapping: dict[str, dict[str, str]] = {}
    for row in countries:
        iso2 = row.get("iso2Code")
        if not iso2:
            continue
        mapping[iso2] = {
            "iso3": row.get("id"),
            "country_name_wb": row.get("name"),
            "region": (row.get("region") or {}).get("value"),
            "income_level": (row.get("incomeLevel") or {}).get("value"),
            "lending_type": (row.get("lendingType") or {}).get("value"),
        }
    mapping.setdefault(
        "XK",
        {
            "iso3": "XKX",
            "country_name_wb": "Kosovo",
            "region": "Europe & Central Asia",
            "income_level": "Upper middle income",
            "lending_type": "IDA",
        },
    )
    return mapping


def metadata_for(country_code: str, country_metadata: dict[str, dict[str, str]]) -> dict[str, str]:
    meta = country_metadata.get(country_code, {})
    return {
        "iso3": meta.get("iso3", ""),
        "country_name_wb": meta.get("country_name_wb", ""),
        "region": meta.get("region", ""),
        "income_level": meta.get("income_level", ""),
        "lending_type": meta.get("lending_type", ""),
        "country_mapping_status": "matched_world_bank_iso2" if meta.get("iso3") else "unmatched_bloomberg_country_code",
    }


def read_static_exports(input_zip: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    rows: list[dict[str, Any]] = []
    inventory: list[dict[str, Any]] = []
    with zipfile.ZipFile(input_zip) as zf:
        for member in sorted(zf.namelist()):
            if not member.lower().endswith((".xlsx", ".xlsm", ".xls")):
                continue
            data = zf.read(member)
            workbook_values = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            workbook_formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            sheet_name = workbook_values.sheetnames[0]
            ws_values = workbook_values[sheet_name]
            ws_formulas = workbook_formulas[sheet_name]
            iterator = ws_values.iter_rows(values_only=True)
            header = [str(x).strip() if x is not None else "" for x in next(iterator)]
            formula_counts: Counter[str] = Counter()
            for formula_row in ws_formulas.iter_rows(min_row=2, values_only=False):
                for idx, cell in enumerate(formula_row):
                    if cell.data_type == "f":
                        formula_counts[header[idx] if idx < len(header) else f"column_{idx + 1}"] += 1

            file_rows = []
            for raw_row in iterator:
                record = {column: raw_row[idx] if idx < len(raw_row) else None for idx, column in enumerate(header)}
                parsed_date = parse_date(record.get("日期"))
                record["_source_file"] = member
                record["_parsed_date"] = parsed_date
                record["_parsed_year"] = parsed_date.year if parsed_date else None
                file_rows.append(record)
                rows.append(record)

            dates = [r["_parsed_date"] for r in file_rows if r["_parsed_date"]]
            inventory.append(
                {
                    "source_file": member,
                    "sheet_name": sheet_name,
                    "row_count": len(file_rows),
                    "column_count": len(header),
                    "date_min": min(dates).isoformat() if dates else "",
                    "date_max": max(dates).isoformat() if dates else "",
                    "formula_cell_count": sum(formula_counts.values()),
                    "formula_columns": collapse_values(list(formula_counts.elements())),
                    "headers": collapse_values(header),
                }
            )
    return rows, inventory


def normalize_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    normalized = []
    for row in rows:
        parsed_date = row.get("_parsed_date")
        if not parsed_date or parsed_date.year < min(YEARS) or parsed_date.year > max(YEARS):
            continue
        agency_raw = str(row.get("机构") or "").strip()
        rating_type_raw = str(row.get("评级类型") or "").strip()
        country_code = str(row.get("国家/地区") or "").strip()
        if not agency_raw or not rating_type_raw or not country_code:
            continue
        normalized.append(
            {
                "source_file": row.get("_source_file"),
                "company_name": str(row.get("公司名称") or "").strip(),
                "event_date": parsed_date.isoformat(),
                "event_year": parsed_date.year,
                "rating_type_raw": rating_type_raw,
                "rating_type": RATING_TYPE_MAP.get(rating_type_raw, rating_type_raw),
                "agency_raw": agency_raw,
                "agency": AGENCY_MAP.get(agency_raw, agency_raw),
                "current_rating_raw": str(row.get("目前评级") or "").strip(),
                "previous_rating_raw": str(row.get("上次评级") or "").strip(),
                "current_rating_clean": clean_rating(row.get("目前评级")) or "",
                "previous_rating_clean": clean_rating(row.get("上次评级")) or "",
                "country_code_bloomberg": country_code,
                "industry_type": str(row.get("行业类型") or "").strip(),
                "security_name": str(row.get("证券名称") or "").strip(),
            }
        )
    return normalized


def exact_dedupe(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    seen = set()
    deduped = []
    key_columns = [
        "company_name",
        "event_date",
        "rating_type_raw",
        "agency_raw",
        "current_rating_raw",
        "previous_rating_raw",
        "country_code_bloomberg",
        "industry_type",
        "security_name",
    ]
    for row in rows:
        key = tuple(row[column] for column in key_columns)
        if key not in seen:
            seen.add(key)
            deduped.append(row)
    return deduped


def group_events(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    grouped: defaultdict[tuple[str, str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        key = (
            row["country_code_bloomberg"],
            row["agency"],
            row["rating_type"],
            row["event_date"],
        )
        grouped[key].append(row)

    events = []
    conflicts = []
    for (country_code, agency, rating_type, event_date), group_rows in sorted(grouped.items()):
        current_values = [row["current_rating_clean"] for row in group_rows if row["current_rating_clean"]]
        previous_values = [row["previous_rating_clean"] for row in group_rows if row["previous_rating_clean"]]
        event = {
            "country_code_bloomberg": country_code,
            "agency": agency,
            "rating_type": rating_type,
            "event_date": event_date,
            "event_year": int(event_date[:4]),
            "current_rating_clean": mode_or_none(current_values) or "",
            "previous_rating_clean": mode_or_none(previous_values) or "",
            "current_rating_raw": mode_or_none([row["current_rating_raw"] for row in group_rows]) or "",
            "previous_rating_raw": mode_or_none([row["previous_rating_raw"] for row in group_rows]) or "",
            "source_row_count": len(group_rows),
            "source_files": collapse_values([row["source_file"] for row in group_rows]),
            "company_names": collapse_values([row["company_name"] for row in group_rows]),
            "security_names": collapse_values([row["security_name"] for row in group_rows]),
            "current_rating_values_all": collapse_values(current_values),
            "previous_rating_values_all": collapse_values(previous_values),
            "current_rating_conflict": len(set(current_values)) > 1,
            "previous_rating_conflict": len(set(previous_values)) > 1,
            "preferred_external_rating_type": (agency, rating_type) in PREFERRED_PAIRS,
        }
        events.append(event)
        if event["current_rating_conflict"] or event["previous_rating_conflict"]:
            conflicts.append(event)
    return events, conflicts


def build_long_panel(events: list[dict[str, Any]], country_codes: list[str], country_metadata: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    by_series: defaultdict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for event in events:
        by_series[(event["country_code_bloomberg"], event["agency"], event["rating_type"])].append(event)

    panel = []
    for country_code in country_codes:
        series_keys = [key for key in by_series if key[0] == country_code]
        for _, agency, rating_type in sorted(series_keys):
            series_events = sorted(by_series[(country_code, agency, rating_type)], key=lambda e: e["event_date"])
            for year in YEARS:
                jan1 = date(year, 1, 1)
                prior = [event for event in series_events if date.fromisoformat(event["event_date"]) <= jan1]
                if prior:
                    source_event = prior[-1]
                    rating_clean = source_event["current_rating_clean"]
                    rating_raw = source_event["current_rating_raw"]
                    source_rule = "latest_event_on_or_before_jan1_current_rating"
                else:
                    after = [event for event in series_events if date.fromisoformat(event["event_date"]) > jan1]
                    if after and date.fromisoformat(after[0]["event_date"]).year == year:
                        source_event = after[0]
                        rating_clean = source_event["previous_rating_clean"]
                        rating_raw = source_event["previous_rating_raw"]
                        source_rule = "first_event_in_year_previous_rating"
                    else:
                        source_event = None
                        rating_clean = ""
                        rating_raw = ""
                        source_rule = "no_prior_or_same_year_event"
                meta = metadata_for(country_code, country_metadata)
                panel.append(
                    {
                        "country_code_bloomberg": country_code,
                        **meta,
                        "year": year,
                        "agency": agency,
                        "rating_type": rating_type,
                        "rating_clean": rating_clean,
                        "rating_raw": rating_raw,
                        "active_rating": is_active_rating(rating_clean),
                        "source_event_date": source_event["event_date"] if source_event else "",
                        "source_rule": source_rule,
                        "source_current_rating_conflict": source_event["current_rating_conflict"] if source_event else False,
                        "source_previous_rating_conflict": source_event["previous_rating_conflict"] if source_event else False,
                        "preferred_external_rating_type": (agency, rating_type) in PREFERRED_PAIRS,
                    }
                )
    return panel


def build_preferred_country_year_panel(long_panel: list[dict[str, Any]], country_codes: list[str], country_metadata: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    preferred_rows = [row for row in long_panel if row["preferred_external_rating_type"]]
    by_country_year_agency: dict[tuple[str, int, str], dict[str, Any]] = {}
    for row in preferred_rows:
        key = (row["country_code_bloomberg"], row["year"], row["agency"])
        if key not in by_country_year_agency or row["active_rating"]:
            by_country_year_agency[key] = row

    out = []
    for country_code in country_codes:
        meta = metadata_for(country_code, country_metadata)
        for year in YEARS:
            moodys = by_country_year_agency.get((country_code, year, "Moodys"))
            fitch = by_country_year_agency.get((country_code, year, "Fitch"))
            sp = by_country_year_agency.get((country_code, year, "SP"))
            active_agencies = [
                agency
                for agency, row in (("Moodys", moodys), ("Fitch", fitch), ("SP", sp))
                if row and row["active_rating"]
            ]
            out.append(
                {
                    "country_code_bloomberg": country_code,
                    **meta,
                    "year": year,
                    "moodys_rating": moodys["rating_clean"] if moodys and moodys["active_rating"] else "",
                    "moodys_source_event_date": moodys["source_event_date"] if moodys and moodys["active_rating"] else "",
                    "moodys_source_rule": moodys["source_rule"] if moodys and moodys["active_rating"] else "",
                    "fitch_rating": fitch["rating_clean"] if fitch and fitch["active_rating"] else "",
                    "fitch_source_event_date": fitch["source_event_date"] if fitch and fitch["active_rating"] else "",
                    "fitch_source_rule": fitch["source_rule"] if fitch and fitch["active_rating"] else "",
                    "sp_rating": sp["rating_clean"] if sp and sp["active_rating"] else "",
                    "sp_source_event_date": sp["source_event_date"] if sp and sp["active_rating"] else "",
                    "sp_source_rule": sp["source_rule"] if sp and sp["active_rating"] else "",
                    "preferred_agency_count": len(active_agencies),
                    "preferred_agencies_available": "; ".join(active_agencies),
                    "has_any_preferred_rating": len(active_agencies) > 0,
                }
            )
    return out


def build_single_precedence_panel(preferred_country_year: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for row in preferred_country_year:
        selected_agency = ""
        selected_rating = ""
        selected_source_event_date = ""
        selected_source_rule = ""
        for agency_key, agency_label in (("moodys", "Moodys"), ("fitch", "Fitch"), ("sp", "SP")):
            rating = row.get(f"{agency_key}_rating", "")
            if rating:
                selected_agency = agency_label
                selected_rating = rating
                selected_source_event_date = row.get(f"{agency_key}_source_event_date", "")
                selected_source_rule = row.get(f"{agency_key}_source_rule", "")
                break
        out.append(
            {
                **row,
                "selected_rating": selected_rating,
                "selected_agency": selected_agency,
                "selected_source_event_date": selected_source_event_date,
                "selected_source_rule": selected_source_rule,
                "has_selected_precedence_rating": bool(selected_rating),
                "agency_precedence_rule": "Moodys issuer_foreign_currency, then Fitch long_term_issuer_default, then SP long_term_foreign_currency_issuer_credit",
            }
        )
    return out


def build_any_country_year_availability(long_panel: list[dict[str, Any]], country_codes: list[str], country_metadata: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    by_country_year: defaultdict[tuple[str, int], list[dict[str, Any]]] = defaultdict(list)
    for row in long_panel:
        if row["active_rating"]:
            by_country_year[(row["country_code_bloomberg"], row["year"])].append(row)
    out = []
    for country_code in country_codes:
        meta = metadata_for(country_code, country_metadata)
        for year in YEARS:
            rows = by_country_year.get((country_code, year), [])
            out.append(
                {
                    "country_code_bloomberg": country_code,
                    **meta,
                    "year": year,
                    "has_any_rating_type": bool(rows),
                    "active_rating_count": len(rows),
                    "active_agencies": collapse_values([row["agency"] for row in rows]),
                    "active_rating_types": collapse_values([row["rating_type"] for row in rows]),
                }
            )
    return out


def build_country_coverage(country_year_panel: list[dict[str, Any]], rating_flag_column: str) -> list[dict[str, Any]]:
    by_country: defaultdict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in country_year_panel:
        by_country[row["country_code_bloomberg"]].append(row)
    coverage = []
    for country_code, rows in sorted(by_country.items()):
        covered_years = sorted(row["year"] for row in rows if row[rating_flag_column])
        missing_years = sorted(set(YEARS) - set(covered_years))
        first = rows[0]
        coverage.append(
            {
                "country_code_bloomberg": country_code,
                "iso3": first.get("iso3", ""),
                "country_name_wb": first.get("country_name_wb", ""),
                "region": first.get("region", ""),
                "income_level": first.get("income_level", ""),
                "lending_type": first.get("lending_type", ""),
                "country_mapping_status": first.get("country_mapping_status", ""),
                "covered_year_count": len(covered_years),
                "missing_year_count": len(missing_years),
                "has_at_least_one_rating_year": len(covered_years) >= 1,
                "has_at_least_one_missing_year": len(missing_years) >= 1,
                "has_five_or_more_missing_years": len(missing_years) >= 5,
                "complete_2000_2025": len(missing_years) == 0,
                "first_covered_year": min(covered_years) if covered_years else "",
                "last_covered_year": max(covered_years) if covered_years else "",
                "missing_years": ";".join(str(y) for y in missing_years),
            }
        )
    return coverage


def build_missing_preferred_country_years(
    preferred_country_year: list[dict[str, Any]],
    any_country_year: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    any_lookup = {
        (row["country_code_bloomberg"], row["year"]): row
        for row in any_country_year
    }
    missing_rows = []
    for row in preferred_country_year:
        if row["has_any_preferred_rating"]:
            continue
        any_row = any_lookup.get((row["country_code_bloomberg"], row["year"]), {})
        has_any_rating_type = bool(any_row.get("has_any_rating_type", False))
        missing_rows.append(
            {
                "country_code_bloomberg": row["country_code_bloomberg"],
                "iso3": row.get("iso3", ""),
                "country_name_wb": row.get("country_name_wb", ""),
                "region": row.get("region", ""),
                "income_level": row.get("income_level", ""),
                "lending_type": row.get("lending_type", ""),
                "country_mapping_status": row.get("country_mapping_status", ""),
                "year": row["year"],
                "missing_preferred_rating": True,
                "has_any_rating_type": has_any_rating_type,
                "active_nonpreferred_agencies": any_row.get("active_agencies", ""),
                "active_nonpreferred_rating_types": any_row.get("active_rating_types", ""),
                "missing_reason": (
                    "only_nonpreferred_rating_type_available"
                    if has_any_rating_type
                    else "no_active_rating_in_bloomberg_event_panel"
                ),
            }
        )

    by_country: defaultdict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in missing_rows:
        by_country[row["country_code_bloomberg"]].append(row)

    summary_rows = []
    for country_code, rows in sorted(by_country.items()):
        first = rows[0]
        missing_years = sorted(row["year"] for row in rows)
        summary_rows.append(
            {
                "country_code_bloomberg": country_code,
                "iso3": first.get("iso3", ""),
                "country_name_wb": first.get("country_name_wb", ""),
                "region": first.get("region", ""),
                "income_level": first.get("income_level", ""),
                "lending_type": first.get("lending_type", ""),
                "country_mapping_status": first.get("country_mapping_status", ""),
                "missing_preferred_year_count": len(missing_years),
                "missing_preferred_years": ";".join(str(year) for year in missing_years),
                "missing_preferred_years_with_nonpreferred_rating_type": sum(
                    row["has_any_rating_type"] for row in rows
                ),
                "missing_preferred_years_with_no_any_rating": sum(
                    not row["has_any_rating_type"] for row in rows
                ),
            }
        )

    return missing_rows, summary_rows


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("")
        return
    fieldnames = list(rows[0].keys())
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def summarize_coverage(label: str, coverage: list[dict[str, Any]]) -> dict[str, Any]:
    covered_counts = [int(row["covered_year_count"]) for row in coverage]
    return {
        "panel": label,
        "country_count": len(coverage),
        "countries_with_at_least_one_rating_year": sum(row["has_at_least_one_rating_year"] for row in coverage),
        "countries_with_no_rating_years": sum(not row["has_at_least_one_rating_year"] for row in coverage),
        "countries_with_at_least_one_missing_year": sum(row["has_at_least_one_missing_year"] for row in coverage),
        "countries_with_five_or_more_missing_years": sum(row["has_five_or_more_missing_years"] for row in coverage),
        "countries_with_complete_2000_2025_panel": sum(row["complete_2000_2025"] for row in coverage),
        "median_covered_years": statistics.median(covered_counts) if covered_counts else "",
    }


def main() -> None:
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    country_metadata = read_country_metadata(args.country_metadata)

    raw_rows, source_inventory = read_static_exports(args.input_zip)
    normalized_rows = normalize_rows(raw_rows)
    deduped_rows = exact_dedupe(normalized_rows)
    events, conflicts = group_events(deduped_rows)

    country_codes = sorted({row["country_code_bloomberg"] for row in deduped_rows})
    for event in events:
        event.update(metadata_for(event["country_code_bloomberg"], country_metadata))
    for row in deduped_rows:
        row.update(metadata_for(row["country_code_bloomberg"], country_metadata))

    long_panel_all = build_long_panel(events, country_codes, country_metadata)
    long_panel_preferred = [row for row in long_panel_all if row["preferred_external_rating_type"]]
    preferred_country_year = build_preferred_country_year_panel(long_panel_all, country_codes, country_metadata)
    single_precedence_country_year = build_single_precedence_panel(preferred_country_year)
    any_country_year = build_any_country_year_availability(long_panel_all, country_codes, country_metadata)

    preferred_coverage = build_country_coverage(preferred_country_year, "has_any_preferred_rating")
    any_coverage = build_country_coverage(any_country_year, "has_any_rating_type")
    missing_preferred_country_years, missing_preferred_country_summary = build_missing_preferred_country_years(
        preferred_country_year,
        any_country_year,
    )

    agency_coverage = []
    for agency in ["Moodys", "Fitch", "SP"]:
        agency_country_year = []
        for country_code in country_codes:
            meta = metadata_for(country_code, country_metadata)
            for year in YEARS:
                rows = [
                    row
                    for row in long_panel_preferred
                    if row["country_code_bloomberg"] == country_code
                    and row["year"] == year
                    and row["agency"] == agency
                    and row["active_rating"]
                ]
                agency_country_year.append(
                    {
                        "country_code_bloomberg": country_code,
                        **meta,
                        "year": year,
                        "has_rating": bool(rows),
                    }
                )
        summary = summarize_coverage(agency, build_country_coverage(agency_country_year, "has_rating"))
        agency_coverage.append(summary)

    coverage_summary = [
        summarize_coverage("preferred_external_rating_any_agency", preferred_coverage),
        summarize_coverage("any_sovereign_rating_type", any_coverage),
        *agency_coverage,
    ]

    source_summary = [
        {
            "metric": "raw_rows_all_years",
            "value": len(raw_rows),
        },
        {
            "metric": "normalized_unique_event_rows_2000_2025",
            "value": len(deduped_rows),
        },
        {
            "metric": "grouped_country_agency_type_date_events",
            "value": len(events),
        },
        {
            "metric": "event_groups_with_rating_conflicts",
            "value": len(conflicts),
        },
        {
            "metric": "preferred_external_rating_event_groups",
            "value": sum(event["preferred_external_rating_type"] for event in events),
        },
        {
            "metric": "bloomberg_country_region_codes",
            "value": len(country_codes),
        },
        {
            "metric": "unmatched_country_region_codes",
            "value": sum(1 for code in country_codes if not country_metadata.get(code, {}).get("iso3")),
        },
    ]

    method_choices = [
        {
            "choice_id": "deduplicate_exact_rows",
            "choice": "Exact duplicate rows across overlapping files are removed using the nine exported Bloomberg columns.",
            "status": "provisional_method_choice",
        },
        {
            "choice_id": "group_same_day_events",
            "choice": "Rows with the same country, agency, rating type, and date are collapsed to one event using the modal current and previous rating; conflicts are written to an audit file.",
            "status": "provisional_method_choice",
        },
        {
            "choice_id": "preferred_external_rating_types",
            "choice": "Preferred external-rating panel uses Moody's issuer foreign-currency rating, Fitch long-term issuer default rating, and S&P long-term foreign-currency issuer credit rating.",
            "status": "provisional_method_choice",
        },
        {
            "choice_id": "beginning_of_year_rule",
            "choice": "Beginning-of-year rating uses the latest event on or before January 1 current rating; if no prior event exists but the first event occurs later in that year, use that event's previous rating.",
            "status": "provisional_method_choice",
        },
        {
            "choice_id": "rating_cleaning",
            "choice": "Bloomberg watch markers such as *+ and *- are removed; Moody's trailing unsolicited marker u is removed for Moody-style symbols; withdrawal/not-rated symbols remain non-active ratings.",
            "status": "provisional_method_choice",
        },
        {
            "choice_id": "no_composite_rating",
            "choice": "No cross-agency composite or median rating is created; the country-year panel keeps agency-specific ratings and availability counts.",
            "status": "accepted_data_build_choice",
        },
        {
            "choice_id": "single_rating_precedence_for_downstream_experiments",
            "choice": "For downstream single-rating experiments, select Moody's issuer foreign-currency rating first, Fitch long-term issuer default rating second, and S&P long-term foreign-currency issuer credit rating third; retain all agency-specific ratings in the wide panel.",
            "status": "accepted_data_build_choice_not_headline_benchmark_approval",
        },
    ]

    write_csv(args.output_dir / "source_file_inventory.csv", source_inventory)
    write_csv(args.output_dir / "source_build_summary.csv", source_summary)
    write_csv(args.output_dir / "method_choices.csv", method_choices)
    write_csv(args.output_dir / "bloomberg_rating_events_deduplicated_2000_2025.csv", deduped_rows)
    write_csv(args.output_dir / "bloomberg_rating_events_grouped_2000_2025.csv", events)
    write_csv(args.output_dir / "bloomberg_rating_event_conflicts_2000_2025.csv", conflicts)
    write_csv(args.output_dir / "bloomberg_boy_ratings_preferred_long_2000_2025.csv", long_panel_preferred)
    write_csv(args.output_dir / "bloomberg_boy_ratings_preferred_country_year_2000_2025.csv", preferred_country_year)
    write_csv(args.output_dir / "bloomberg_boy_ratings_single_precedence_country_year_2000_2025.csv", single_precedence_country_year)
    write_csv(args.output_dir / "bloomberg_boy_ratings_any_type_country_year_2000_2025.csv", any_country_year)
    write_csv(args.output_dir / "bloomberg_preferred_country_coverage_2000_2025.csv", preferred_coverage)
    write_csv(args.output_dir / "bloomberg_any_type_country_coverage_2000_2025.csv", any_coverage)
    write_csv(args.output_dir / "bloomberg_missing_preferred_country_years_2000_2025.csv", missing_preferred_country_years)
    write_csv(args.output_dir / "bloomberg_missing_preferred_country_summary_2000_2025.csv", missing_preferred_country_summary)
    write_csv(args.output_dir / "coverage_summary.csv", coverage_summary)

    readme = [
        "# Bloomberg Sovereign Rating Panel",
        "",
        "Status: derived from static-value Bloomberg rating-change exports; provisional, not canonical.",
        "",
        "Build script: `scripts/build_bloomberg_sovereign_rating_panel.py`",
        f"Input ZIP: `{args.input_zip}`",
        "",
        "Main panel files:",
        "",
        "- `bloomberg_boy_ratings_preferred_country_year_2000_2025.csv`: one row per country-year with Moody's, Fitch, and S&P preferred external-rating columns.",
        "- `bloomberg_boy_ratings_single_precedence_country_year_2000_2025.csv`: one row per country-year with a selected rating using Moody's, then Fitch, then S&P precedence.",
        "- `bloomberg_boy_ratings_preferred_long_2000_2025.csv`: long agency-specific beginning-of-year ratings.",
        "- `bloomberg_rating_events_grouped_2000_2025.csv`: grouped rating-change events used for carry-forward.",
        "- `bloomberg_rating_event_conflicts_2000_2025.csv`: same-day country-agency-type events with conflicting rating values.",
        "- `bloomberg_missing_preferred_country_years_2000_2025.csv`: explicit country-year ledger for missing preferred external ratings.",
        "- `bloomberg_missing_preferred_country_summary_2000_2025.csv`: country-level summary of missing preferred external-rating years.",
        "- `coverage_summary.csv`: high-level coverage counts.",
        "",
        "The panel is inferred from rating-change events, not directly exported as annual as-of ratings.",
    ]
    (args.output_dir / "README.md").write_text("\n".join(readme) + "\n", encoding="utf-8")

    print(f"Wrote Bloomberg sovereign rating panel to {args.output_dir}")
    for row in coverage_summary:
        print(row)


if __name__ == "__main__":
    main()
